import csv
import io

import pdfplumber
from docx import Document as new_docx

from app.glossary import Glossary, build_glossary
from app.lang import detect_language
from app.pipeline import JobResult, inspect_upload, report_csv, run_job
from app.prompt import BODY_MARKER

EMPTY = Glossary("de", "en", ())
GERMAN = "Der schnelle braune Fuchs springt über den faulen Hund am Flussufer."


def upper(prompt: str) -> str:
    return prompt.split(BODY_MARKER, 1)[1].upper()


def _docx_bytes() -> bytes:
    d = new_docx()
    d.core_properties.author = "Autorin"
    d.add_heading("Kapitel", 1)
    d.add_paragraph("Ein Absatz")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _job(name: str, data: bytes, glossary: Glossary = EMPTY, llm=upper) -> JobResult:
    return run_job(name, data, glossary, llm, user="T. Hein")


def test_markdown_input_gives_three_consistent_outputs() -> None:
    result = _job("vertrag.md", b"# Titel\n\nEin Absatz\n")
    assert result.stem == "vertrag"
    assert result.target == "en"
    assert result.md == b"# TITEL\n\nEIN ABSATZ\n"
    docx = new_docx(io.BytesIO(result.docx))
    assert [p.text for p in docx.paragraphs] == ["TITEL", "EIN ABSATZ"]
    props = docx.core_properties
    assert (props.title, props.language, props.author) == ("vertrag", "en", "T. Hein")
    pdf = pdfplumber.open(io.BytesIO(result.pdf))
    assert "EIN ABSATZ" in pdf.pages[0].extract_text()


def test_docx_input_is_translated_in_place() -> None:
    result = _job("bericht.docx", _docx_bytes())
    docx = new_docx(io.BytesIO(result.docx))
    assert [p.text for p in docx.paragraphs] == ["KAPITEL", "EIN ABSATZ"]
    assert docx.core_properties.author == "Autorin"  # kept, not replaced by the user
    assert docx.core_properties.language == "en"
    assert b"EIN ABSATZ" in result.md


def test_report_lists_terms_failed_blocks_and_untouched_features() -> None:
    glossary = build_glossary([["de", "en"], ["Freigabe", "clearance"]], "de", "en").glossary
    result = _job("a.md", b"Die Freigabe\n\nZweiter Absatz\n", glossary)
    assert result.hits == 1
    assert result.enforced == 0
    rows = list(csv.reader(io.StringIO(report_csv(result).decode("utf-8-sig"))))
    assert rows[0] == ["kind", "detail", "source", "expected"]
    assert ["term", "Die Freigabe", "Freigabe", "clearance"] in rows

    def broken(_prompt: str) -> str:
        return "garbage"

    failed = _job("a.md", b"Bleibt\n", llm=broken)
    assert ["failed", "Bleibt", "", ""] in list(
        csv.reader(io.StringIO(report_csv(failed).decode("utf-8-sig")))
    )
    noted = JobResult(**{**vars(result), "untouched": ("comments",)})
    assert ["untouched", "comments", "", ""] in list(
        csv.reader(io.StringIO(report_csv(noted).decode("utf-8-sig")))
    )


def test_cancel_marks_the_result() -> None:
    body = "\n\n".join(f"Absatz {i}" for i in range(5))
    result = run_job(
        "a.md",
        body.encode(),
        EMPTY,
        upper,
        user="x",
        cancel=lambda: True,
        segment_chars=20,
    )
    assert result.cancelled


def test_inspect_upload_detects_language_and_reports_errors() -> None:
    info = inspect_upload("a.txt", GERMAN.encode())
    assert info.language == "de"
    assert info.error == ""
    assert inspect_upload("a.xyz", b"x").error != ""
    assert inspect_upload("a.txt", b"   ").error != ""


def test_detect_language_only_returns_supported_codes() -> None:
    assert detect_language(GERMAN) == "de"
    assert detect_language("12345 67890") is None
