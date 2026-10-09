import io

import pytest
from docx import Document as new_docx
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from PIL import Image

from app.document import Block, EmptyDocumentError, Kind
from app.readers import DocumentReadError, read_document


def _bytes(doc: object) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)  # type: ignore[attr-defined]
    return buf.getvalue()


def _outline(element: object, level: int) -> None:
    ppr = element.get_or_add_pPr()  # type: ignore[attr-defined]
    node = OxmlElement("w:outlineLvl")
    node.set(qn("w:val"), str(level))
    ppr.append(node)


def _read(doc: object) -> list[Block]:
    return read_document("x.docx", _bytes(doc))


def test_title_and_heading_styles_map_to_levels() -> None:
    d = new_docx()
    d.add_heading("Titel", 0)
    d.add_heading("Eins", 1)
    d.add_heading("Zwei", 2)
    d.add_heading("Drei", 3)
    d.add_paragraph("Text")
    blocks = _read(d)
    assert [(b.kind, b.level, b.text) for b in blocks] == [
        (Kind.HEADING, 1, "Titel"),
        (Kind.HEADING, 1, "Eins"),
        (Kind.HEADING, 2, "Zwei"),
        (Kind.HEADING, 3, "Drei"),
        (Kind.PARAGRAPH, 0, "Text"),
    ]


def test_custom_styles_map_through_outline_level() -> None:
    d = new_docx()
    style = d.styles.add_style("Überschrift Eins", WD_STYLE_TYPE.PARAGRAPH)
    _outline(style.element, 0)
    d.add_paragraph("Kapitel", style="Überschrift Eins")
    para = d.add_paragraph("Direkt")
    _outline(para._p, 2)
    blocks = _read(d)
    assert [(b.kind, b.level) for b in blocks] == [(Kind.HEADING, 1), (Kind.HEADING, 3)]


def test_lists_keep_depth_and_order_type() -> None:
    d = new_docx()
    d.add_paragraph("a", style="List Bullet")
    d.add_paragraph("b", style="List Bullet 2")
    d.add_paragraph("c", style="List Number")
    blocks = _read(d)
    assert [(b.kind, b.depth, b.ordered) for b in blocks] == [
        (Kind.LIST_ITEM, 0, False),
        (Kind.LIST_ITEM, 1, False),
        (Kind.LIST_ITEM, 0, True),
    ]


def test_table_keeps_shape_and_cell_text() -> None:
    d = new_docx()
    table = d.add_table(rows=2, cols=3)
    for r, row in enumerate(table.rows):
        for c, cell in enumerate(row.cells):
            cell.text = f"r{r}c{c}"
    (block,) = _read(d)
    assert block.kind is Kind.TABLE
    assert [len(r) for r in block.rows] == [3, 3]
    assert block.rows[1][2] == "r1c2"


def test_bold_italic_runs_become_inline_markup() -> None:
    d = new_docx()
    p = d.add_paragraph("normal ")
    p.add_run("fett ").bold = True
    p.add_run("und").bold = True
    p.add_run(" ")
    p.add_run("kursiv").italic = True
    assert _read(d)[0].text == "normal **fett und** *kursiv*"


def test_hyperlink_becomes_markdown_link() -> None:
    d = new_docx()
    p = d.add_paragraph("Siehe ")
    rid = d.part.relate_to("https://example.com", RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "Seite"
    run.append(text)
    link.append(run)
    p._p.append(link)
    assert Paragraph(p._p, d).text == "Siehe Seite"
    assert _read(d)[0].text == "Siehe [Seite](https://example.com)"


def test_image_becomes_placeholder_and_empty_paragraphs_are_skipped() -> None:
    img = io.BytesIO()
    Image.new("RGB", (4, 4)).save(img, "PNG")
    d = new_docx()
    d.add_paragraph("")
    d.add_picture(io.BytesIO(img.getvalue()))
    d.add_paragraph("Text")
    blocks = _read(d)
    assert [b.kind for b in blocks] == [Kind.IMAGE_PLACEHOLDER, Kind.PARAGRAPH]
    assert blocks[0].text.startswith("[Bild: ")
    assert not blocks[0].translate


def test_empty_docx_raises_empty_error() -> None:
    with pytest.raises(EmptyDocumentError):
        _read(new_docx())


def test_corrupt_docx_raises_read_error() -> None:
    with pytest.raises(DocumentReadError, match="DOCX"):
        read_document("x.docx", b"not a zip")


def test_md_txt_and_dispatch_errors() -> None:
    assert read_document("a.MD", b"# T\n\ntext")[0].kind is Kind.HEADING
    blocks = read_document("a.txt", b"Zeile eins\nZeile zwei\n\nAbsatz zwei")
    assert [b.text for b in blocks] == ["Zeile eins Zeile zwei", "Absatz zwei"]
    with pytest.raises(DocumentReadError, match="Dateityp"):
        read_document("a.rtf", b"x")
    with pytest.raises(EmptyDocumentError):
        read_document("a.txt", b"  \n ")
