import io

import pdfplumber
import pdfplumber.pdf
import pytest
from docx import Document as load_docx
from docx.enum.section import WD_ORIENT

from app.document import Block, Document, Kind, parse_markdown, same_structure, write_markdown
from app.readers import read_document
from app.writers import mime_type, output_name, write_docx, write_md, write_pdf

SAMPLE = """# Titel

Ein Absatz mit **fett**, *kursiv* und [Link](https://a.example).

## Liste

- eins
  - zwei
- drei

1. erst
2. dann

| A | B |
|---|---|
| 1 | 2 |
| 3 | 4 |
"""


def _docx(doc: Document) -> object:
    return load_docx(io.BytesIO(write_docx(doc, title="vertrag", language="de", author="Hein")))


def _pdf(doc: Document, **kw: str) -> pdfplumber.pdf.PDF:
    data = write_pdf(doc, title=kw.get("title", "t"), author="Hein")
    assert data.startswith(b"%PDF")
    return pdfplumber.open(io.BytesIO(data))


def test_output_name_and_mime_types() -> None:
    assert output_name("vertrag", "de", "docx") == "vertrag_de.docx"
    assert mime_type("md") == "text/markdown"
    assert mime_type("pdf") == "application/pdf"
    assert mime_type("docx").endswith("wordprocessingml.document")


def test_md_equals_m1_writer() -> None:
    doc = parse_markdown(SAMPLE)
    assert write_md(doc) == write_markdown(doc).encode()


def test_docx_metadata_and_styles() -> None:
    d = _docx(parse_markdown(SAMPLE))
    props = d.core_properties  # type: ignore[attr-defined]
    assert (props.title, props.language, props.author) == ("vertrag", "de", "Hein")
    styles = [p.style.name for p in d.paragraphs]  # type: ignore[attr-defined]
    assert styles[0] == "Heading 1"
    assert "Heading 2" in styles
    assert "List Bullet 2" in styles
    assert "List Number" in styles
    assert len(d.tables) == 1  # type: ignore[attr-defined]


def test_docx_runs_and_hyperlink() -> None:
    d = _docx(parse_markdown(SAMPLE))
    para = d.paragraphs[1]  # type: ignore[attr-defined]
    assert [r.text for r in para.runs if r.bold] == ["fett"]
    assert [r.text for r in para.runs if r.italic] == ["kursiv"]
    assert [(h.text, h.url) for h in para.hyperlinks] == [("Link", "https://a.example")]


def test_docx_round_trip_keeps_structure() -> None:
    doc = parse_markdown(SAMPLE)
    back = read_document("x.docx", write_docx(doc, title="t", language="de", author="a"))
    assert same_structure(doc, back)
    assert back[1].text == doc[1].text


def test_docx_other_kinds_do_not_fail() -> None:
    doc = [
        Block(Kind.QUOTE, "Zitat\nzweite Zeile"),
        Block(Kind.CODE, "x = 1", translate=False),
        Block(Kind.IMAGE_PLACEHOLDER, "[Bild: a]", translate=False),
        Block(Kind.PAGE_BREAK, translate=False),
        Block(Kind.LIST_ITEM, "tief", depth=7),
    ]
    text = "\n".join(p.text for p in _docx(doc).paragraphs)  # type: ignore[attr-defined]
    assert "Zitat" in text
    assert "x = 1" in text
    assert "[Bild: a]" in text


def test_docx_wide_table_is_landscape() -> None:
    wide = Block(Kind.TABLE, rows=(tuple("abcdefghi"), tuple("123456789")))
    section = _docx([wide]).sections[0]  # type: ignore[attr-defined]
    assert section.orientation == WD_ORIENT.LANDSCAPE
    assert section.page_width > section.page_height


def test_pdf_text_and_special_characters() -> None:
    text = "Zürich, Café, Łódź, Čeština, ß"
    pdf = _pdf([Block(Kind.PARAGRAPH, text)])
    assert text in pdf.pages[0].extract_text()


def test_pdf_headings_are_larger_and_bold() -> None:
    pdf = _pdf(parse_markdown(SAMPLE))
    chars = pdf.pages[0].chars
    head = next(c for c in chars if c["text"] == "T")
    body = next(c for c in chars if c["text"] == "E")
    assert head["size"] > body["size"]
    assert "Bold" in head["fontname"]
    assert "Bold" not in body["fontname"]


def test_pdf_draws_table_and_lists() -> None:
    page = _pdf(parse_markdown(SAMPLE)).pages[0]
    assert page.extract_tables()[0] == [["A", "B"], ["1", "2"], ["3", "4"]]
    text = page.extract_text()
    assert "• eins" in text
    assert "2. dann" in text


def test_pdf_long_table_breaks_across_pages() -> None:
    rows = tuple((f"z{i}", "x") for i in range(150))
    pdf = _pdf([Block(Kind.TABLE, rows=(("K", "V"), *rows))])
    assert len(pdf.pages) > 1
    assert "K V" in pdf.pages[1].extract_text()  # header row repeats


def test_pdf_wide_table_is_landscape() -> None:
    wide = Block(Kind.TABLE, rows=(tuple("abcdefghi"), tuple("123456789")))
    page = _pdf([wide]).pages[0]
    assert page.width > page.height


@pytest.mark.parametrize("fillers", range(36, 52))
def test_pdf_heading_stays_with_next_block(fillers: int) -> None:
    doc = [Block(Kind.PARAGRAPH, f"Füller {i}") for i in range(fillers)]
    doc += [Block(Kind.HEADING, "Kopf", level=2), Block(Kind.PARAGRAPH, "Danach")]
    pdf = _pdf(doc)
    where = {
        w: p.page_number for p in pdf.pages for w in ("Kopf", "Danach") if w in p.extract_text()
    }
    assert where["Kopf"] == where["Danach"]


def test_glossary_template_as_xlsx_reads_like_the_csv() -> None:
    from app.glossary import GLOSSARY_TEMPLATE_CSV
    from app.readers import read_glossary_rows
    from app.writers import glossary_template_xlsx

    xlsx = read_glossary_rows("vorlage.xlsx", glossary_template_xlsx())
    assert xlsx == read_glossary_rows("vorlage.csv", GLOSSARY_TEMPLATE_CSV.encode())
