import io
import re
import zipfile
from typing import Any

import pdfplumber
import pytest
from docx import Document as new_docx
from docx.enum.text import WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt
from PIL import Image

from app.document import DocumentReadError, parse_markdown
from app.docx_inplace import translate_docx
from app.glossary import Glossary, build_glossary
from app.prompt import BODY_MARKER
from app.writers import write_docx, write_md, write_pdf

EMPTY = Glossary("de", "en", ())
TAG = re.compile(r"(</?(?:a\d+|b|i|u)>)")
EN = "The quick brown fox jumps over the lazy dog near the river bank."


class Stub:
    """Upper-cases the text but keeps tags; records every prompt."""

    def __init__(self, keep_tags: bool = True) -> None:
        self.keep_tags = keep_tags
        self.prompts: list[str] = []

    def __call__(self, prompt: str) -> str:
        self.prompts.append(prompt)
        body = prompt.split(BODY_MARKER, 1)[1]
        if not self.keep_tags:
            body = TAG.sub("", body)
        return "".join(p if TAG.fullmatch(p) else p.upper() for p in TAG.split(body))

    @property
    def sent(self) -> str:
        return "\n".join(self.prompts)


def _save(doc: Any) -> bytes:
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _run(data: bytes, stub: Stub | None = None, glossary: Glossary = EMPTY) -> Any:
    return translate_docx(data, glossary, stub or Stub())


def _out(result: Any) -> Any:
    return new_docx(io.BytesIO(result.data))


def _png(tmp_path: Any) -> str:
    path = tmp_path / "p.png"
    Image.new("RGB", (4, 4)).save(path)
    return str(path)


def test_body_tables_nested_tables_headers_and_footers_are_translated() -> None:
    d = new_docx()
    d.add_paragraph("Erster Absatz")
    table = d.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    cell.paragraphs[0].text = "Zelle"
    cell.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0].text = "Innere Zelle"
    section = d.sections[0]
    section.header.is_linked_to_previous = False
    section.header.paragraphs[0].text = "Kopfzeile"
    section.footer.is_linked_to_previous = False
    section.footer.paragraphs[0].text = "Fusszeile"
    out = _out(_run(_save(d)))
    assert out.paragraphs[0].text == "ERSTER ABSATZ"
    assert out.tables[0].cell(0, 0).paragraphs[0].text == "ZELLE"
    assert out.tables[0].cell(0, 0).tables[0].cell(0, 0).text == "INNERE ZELLE"
    assert out.sections[0].header.paragraphs[0].text == "KOPFZEILE"
    assert out.sections[0].footer.paragraphs[0].text == "FUSSZEILE"


def test_footnotes_comments_and_text_boxes_stay_and_are_reported() -> None:
    d = new_docx()
    para = d.add_paragraph("Haupttext")
    d.add_comment(para.runs, text="Kommentartext", author="X")
    box = parse_xml(
        f'<w:r {nsdecls("w")} xmlns:v="urn:schemas-microsoft-com:vml"><w:pict><v:shape>'
        "<v:textbox><w:txbxContent><w:p><w:r><w:t>Im Kasten</w:t></w:r></w:p></w:txbxContent>"
        "</v:textbox></v:shape></w:pict></w:r>"
    )
    para._p.append(box)
    foot = (
        f'<w:footnotes {nsdecls("w")}><w:footnote w:id="1"><w:p><w:r><w:t>Fußnote</w:t></w:r></w:p>'
        "</w:footnote></w:footnotes>"
    )
    part = Part(
        PackURI("/word/footnotes.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
        foot.encode(),
        d.part.package,
    )
    d.part.relate_to(part, RT.FOOTNOTES)
    stub = Stub()
    result = _run(_save(d), stub)
    assert set(result.untouched) == {"footnotes", "comments", "text boxes"}
    assert "Im Kasten" not in stub.sent
    assert "Kommentartext" not in stub.sent
    with zipfile.ZipFile(io.BytesIO(result.data)) as z:
        assert "Fußnote" in z.read("word/footnotes.xml").decode()
        assert "Im Kasten" in z.read("word/document.xml").decode()
        assert "Kommentartext" in z.read("word/comments.xml").decode()


def test_styles_numbering_images_sections_and_properties_are_unchanged(tmp_path: Any) -> None:
    d = new_docx()
    d.add_heading("Überschrift", 1)
    d.add_paragraph("Punkt", style="List Bullet")
    d.add_picture(_png(tmp_path))
    d.core_properties.author = "Autorin"
    d.core_properties.title = "Titel"
    d.core_properties.language = "de"
    src = _save(d)
    out = _run(src).data
    before, after = zipfile.ZipFile(io.BytesIO(src)), zipfile.ZipFile(io.BytesIO(out))
    for name in before.namelist():
        if name.startswith("word/media/") or name in ("word/styles.xml", "word/numbering.xml"):
            assert before.read(name) == after.read(name), name
    src_doc, out_doc = new_docx(io.BytesIO(src)), new_docx(io.BytesIO(out))
    for a, b in zip(src_doc.paragraphs, out_doc.paragraphs, strict=True):
        assert a.style is not None
        assert b.style is not None
        assert a.style.name == b.style.name
        assert (a._p.pPr is None) == (b._p.pPr is None)
    assert src_doc.sections[0]._sectPr.xml == out_doc.sections[0]._sectPr.xml
    props = out_doc.core_properties
    assert (props.author, props.title, props.language) == ("Autorin", "Titel", "en")
    assert len(out_doc.inline_shapes) == 1


def test_bold_and_italic_survive_with_valid_tags() -> None:
    d = new_docx()
    p = d.add_paragraph()
    p.add_run("normal ")
    bold = p.add_run("bold")
    bold.bold = True
    bold.font.size = Pt(14)
    p.add_run(" normal ")
    p.add_run("italic").italic = True
    runs = _out(_run(_save(d))).paragraphs[0].runs
    assert [r.text for r in runs if r.bold] == ["BOLD"]
    assert [r.text for r in runs if r.italic] == ["ITALIC"]
    assert next(r for r in runs if r.bold).font.size == Pt(14)
    assert "".join(r.text for r in runs) == "NORMAL BOLD NORMAL ITALIC"


def test_broken_tags_fall_back_to_one_run_with_first_run_formatting() -> None:
    d = new_docx()
    p = d.add_paragraph()
    p.add_run("normal ").font.name = "Arial"
    p.add_run("bold").bold = True
    runs = _out(_run(_save(d), Stub(keep_tags=False))).paragraphs[0].runs
    assert len(runs) == 1
    assert runs[0].text == "NORMAL BOLD"
    assert runs[0].font.name == "Arial"
    assert not runs[0].bold


def test_glossary_miss_highlights_all_runs_of_that_paragraph_only() -> None:
    glossary = build_glossary([["de", "en"], ["Freigabe", "clearance"]], "de", "en").glossary
    d = new_docx()
    p = d.add_paragraph("Die ")
    p.add_run("Freigabe").bold = True
    d.add_paragraph("Ohne Begriff")
    out = _out(_run(_save(d), glossary=glossary))
    assert [r.font.highlight_color for r in out.paragraphs[0].runs] == [WD_COLOR_INDEX.YELLOW] * 2
    assert all(r.font.highlight_color is None for r in out.paragraphs[1].runs)


def test_text_already_in_target_language_is_not_sent() -> None:
    d = new_docx()
    d.add_paragraph(EN)
    d.add_paragraph("Ein deutscher Absatz")
    stub = Stub()
    out = _out(_run(_save(d), stub))
    assert out.paragraphs[0].text == EN
    assert out.paragraphs[1].text == "EIN DEUTSCHER ABSATZ"
    assert "quick brown fox" not in stub.sent


def test_empty_image_only_and_field_paragraphs_are_left_alone(tmp_path: Any) -> None:
    d = new_docx()
    d.add_paragraph("")
    d.add_paragraph().add_run().add_picture(_png(tmp_path))
    field = parse_xml(
        f'<w:p {nsdecls("w")}><w:r><w:fldChar w:fldCharType="begin"/></w:r>'
        '<w:r><w:instrText> PAGE </w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate"/>'
        '</w:r><w:r><w:t>Seite drei</w:t></w:r><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p>'
    )
    d.element.body.insert(0, field)
    d.add_paragraph("Text")
    stub = Stub()
    out = _out(_run(_save(d), stub))
    assert "Seite drei" not in stub.sent
    assert "Seite drei" in out.element.xml
    assert len(out.inline_shapes) == 1


def test_hyperlink_keeps_target_and_translates_text() -> None:
    src = write_docx(
        parse_markdown("Siehe [Seite](https://a.example) hier"),
        title="t",
        language="de",
        author="a",
    )
    out = _out(_run(src))
    para = out.paragraphs[0]
    assert [(h.text, h.url) for h in para.hyperlinks] == [("SEITE", "https://a.example")]
    assert para.text == "SIEHE SEITE HIER"


def test_md_and_pdf_carry_the_same_translated_text() -> None:
    d = new_docx()
    d.add_heading("Kapitel", 1)
    d.add_paragraph("Ein Absatz")
    result = _run(_save(d))
    assert "EIN ABSATZ" in write_md(result.document).decode()
    pdf = pdfplumber.open(io.BytesIO(write_pdf(result.document, title="t", author="a")))
    assert "EIN ABSATZ" in pdf.pages[0].extract_text()
    assert "KAPITEL" in pdf.pages[0].extract_text()


def test_failed_paragraphs_keep_the_source_text() -> None:
    d = new_docx()
    d.add_paragraph("Bleibt")

    def broken(_prompt: str) -> str:
        return "garbage"

    out = _out(translate_docx(_save(d), EMPTY, broken))
    assert out.paragraphs[0].text == "Bleibt"


def test_corrupt_input_raises_read_error() -> None:
    with pytest.raises(DocumentReadError):
        _run(b"not a zip")
