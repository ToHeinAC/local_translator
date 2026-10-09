"""DOCX to Document adapter (python-docx). Headers, footers and notes are not read."""

# python-docx exposes its lxml elements (`_p`, `_r`, `xpath`) untyped; this adapter is the one
# place that touches them.
# pyright: reportPrivateUsage=false, reportUnknownMemberType=false, reportUnknownVariableType=false

import io
import re
import zipfile
from collections.abc import Iterator
from itertools import groupby
from typing import Any

from docx import Document as load_docx
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.hyperlink import Hyperlink
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from app.document import Block, Document, DocumentReadError, Kind, require_text

_Piece = tuple[str, bool, bool]  # text, bold, italic
_HEADING = re.compile(r"Heading (\d)")
_LIST_DEPTH = re.compile(r"List (?:Bullet|Number|Paragraph) (\d)")


def read_docx(data: bytes) -> Document:
    """Body paragraphs and tables in document order; raises ``DocumentReadError`` if corrupt."""
    try:
        doc = load_docx(io.BytesIO(data))
    except (PackageNotFoundError, zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise DocumentReadError("Ungültige oder beschädigte DOCX-Datei") from exc
    numbering = _numbering_root(doc)
    blocks: Document = []
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            blocks += _paragraph(Paragraph(child, doc), numbering)  # type: ignore[arg-type]
        elif child.tag == qn("w:tbl"):
            blocks.append(_table(Table(child, doc)))  # type: ignore[arg-type]
    return require_text(blocks)


def _numbering_root(doc: object) -> object | None:
    try:
        return doc.part.numbering_part.element  # type: ignore[attr-defined]
    except (NotImplementedError, KeyError, AttributeError):
        return None


def _paragraph(par: Paragraph, numbering: object | None) -> list[Block]:
    text, images = _inline(par)
    placeholders = [Block(Kind.IMAGE_PLACEHOLDER, f"[Bild: {n}]", translate=False) for n in images]
    if not text.strip():
        return placeholders
    return [_text_block(par, text, numbering), *placeholders]


def _text_block(par: Paragraph, text: str, numbering: object | None) -> Block:
    level = _heading_level(par)
    if level:
        return Block(Kind.HEADING, text, level=level)
    item = _list_info(par, numbering)
    if item:
        return Block(Kind.LIST_ITEM, text, depth=item[0], ordered=item[1])
    return Block(Kind.PARAGRAPH, text)


def _styles(par: Paragraph) -> Iterator[Any]:
    style: Any = par.style
    while style is not None:
        yield style
        style = style.base_style


def _vals(element: object, path: str) -> list[str]:
    return [str(v) for v in element.xpath(path)]  # type: ignore[attr-defined]


def _heading_level(par: Paragraph) -> int:
    name = par.style.name if par.style is not None else ""
    if name == "Title":
        return 1
    if m := _HEADING.fullmatch(name or ""):
        return min(int(m.group(1)), 6)
    outline = _vals(par._p, "./w:pPr/w:outlineLvl/@w:val")
    for style in _styles(par):
        outline = outline or _vals(style.element, "./w:pPr/w:outlineLvl/@w:val")  # type: ignore[attr-defined]
    if outline and int(outline[0]) < 9:
        return min(int(outline[0]) + 1, 6)
    return 0


def _num_pr(par: Paragraph) -> object | None:
    found = par._p.xpath("./w:pPr/w:numPr")
    for style in _styles(par):
        found = found or style.element.xpath("./w:pPr/w:numPr")  # type: ignore[attr-defined]
    return found[0] if found else None


def _list_info(par: Paragraph, numbering: object | None) -> tuple[int, bool] | None:
    num_pr = _num_pr(par)
    num_id = _vals(num_pr, "./w:numId/@w:val") if num_pr is not None else []
    if not num_id or num_id[0] == "0":
        return None
    ilvl = _vals(num_pr, "./w:ilvl/@w:val")
    name = par.style.name if par.style is not None else ""
    suffix = _LIST_DEPTH.fullmatch(name or "")
    depth = int(ilvl[0]) if ilvl else int(suffix.group(1)) - 1 if suffix else 0
    fmt = _num_format(numbering, num_id[0], depth)
    return depth, (fmt != "bullet") if fmt else "Number" in (name or "")


def _num_format(numbering: object | None, num_id: str, ilvl: int) -> str:
    if numbering is None:
        return ""
    abstract = _vals(numbering, f'w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')
    if not abstract:
        return ""
    path = (
        f'w:abstractNum[@w:abstractNumId="{abstract[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val'
    )
    return next(iter(_vals(numbering, path)), "")


def _image_name(run: Run) -> str | None:
    if not run._r.xpath(".//w:drawing | .//w:pict"):
        return None
    names = _vals(run._r, ".//wp:docPr/@descr") + _vals(run._r, ".//wp:docPr/@name")
    return next((n for n in names if n), "Bild")


def _inline(par: Paragraph) -> tuple[str, list[str]]:
    """Paragraph text in Markdown inline syntax, plus the names of images it contains."""
    pieces: list[_Piece] = []
    images: list[str] = []
    for item in par.iter_inner_content():
        if isinstance(item, Hyperlink):
            pieces.append((f"[{item.text}]({item.url})" if item.url else item.text, False, False))
        elif (image := _image_name(item)) is not None:
            images.append(image)
        else:
            pieces.append((item.text, bool(item.bold), bool(item.italic)))
    merged = [
        ("".join(p[0] for p in group), fmt[0], fmt[1])
        for fmt, group in groupby(pieces, key=lambda p: (p[1], p[2]))
    ]
    return "".join(_wrap(*p) for p in merged).strip(), images


def _wrap(text: str, bold: bool, italic: bool) -> str:
    core = text.strip()
    if not core or not (bold or italic):
        return text
    mark = "*" * ((2 if bold else 0) + (1 if italic else 0))
    lead = text[: len(text) - len(text.lstrip())]
    trail = text[len(text.rstrip()) :]
    return f"{lead}{mark}{core}{mark}{trail}"


def _table(table: Table) -> Block:
    rows = tuple(
        tuple(" ".join(t for p in cell.paragraphs if (t := _inline(p)[0])) for cell in row.cells)
        for row in table.rows
    )
    return Block(Kind.TABLE, rows=rows)
