"""Input adapters: raw file bytes to structured data."""

import csv
import io
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.document import EmptyDocumentError, Kind, parse_markdown
from app.glossary import GlossaryError


def read_glossary_rows(name: str, data: bytes) -> list[list[str]]:
    """Read a csv/tsv/xlsx/md glossary into rows of stripped strings (first row = header)."""
    ext = Path(name).suffix.lower()
    if ext in (".csv", ".tsv"):
        return _read_delimited(data)
    if ext == ".xlsx":
        return _read_xlsx(data)
    if ext == ".md":
        return _read_md_table(data)
    raise GlossaryError(f"Dateityp '{ext}' wird für Glossare nicht unterstützt")


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _read_delimited(data: bytes) -> list[list[str]]:
    text = _decode(data)
    first = text.splitlines()[0] if text.strip() else ""
    delimiter = max(",;\t", key=first.count)
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    return [[c.strip() for c in row] for row in reader]


def _read_xlsx(data: bytes) -> list[list[str]]:
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, InvalidFileException, KeyError) as exc:
        raise GlossaryError("Ungültige xlsx-Datei") from exc
    sheet = book.worksheets[0]
    return [
        ["" if c is None else str(c).strip() for c in row]
        for row in sheet.iter_rows(values_only=True)
    ]


def _read_md_table(data: bytes) -> list[list[str]]:
    try:
        doc = parse_markdown(_decode(data))
    except EmptyDocumentError as exc:
        raise GlossaryError("Keine Tabelle im Markdown-Glossar gefunden") from exc
    for block in doc:
        if block.kind is Kind.TABLE:
            return [[c.strip() for c in row] for row in block.rows]
    raise GlossaryError("Keine Tabelle im Markdown-Glossar gefunden")
