"""Document to DOCX bytes (python-docx, Word's default template)."""

# python-docx ships partial type information; this adapter is one of two places that touch it.
# pyright: reportUnknownMemberType=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false

import io
from typing import Any

from docx import Document as new_docx
from docx.enum.section import WD_ORIENT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor
from docx.text.paragraph import Paragraph
from docx.text.run import Run as DocxRun

from app.document import Block, Document, Kind
from app.inline import Run, parse_inline

WIDE_TABLE_COLUMNS = 8
_MAX_LIST_LEVEL = 3
_CODE_FONT = "Courier New"


def write_docx(doc: Document, *, title: str, language: str, author: str) -> bytes:
    """Rebuild ``doc`` as DOCX; headings and lists use Word's built-in styles."""
    out: Any = new_docx()
    props = out.core_properties
    props.title, props.language, props.author = title, language, author
    if any(b.kind is Kind.TABLE and len(b.rows[0]) > WIDE_TABLE_COLUMNS for b in doc):
        _landscape(out.sections[0])
    for block in doc:
        _block(out, block)
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _landscape(section: Any) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width


def _block(out: Any, block: Block) -> None:
    match block.kind:
        case Kind.HEADING:
            _runs(out.add_paragraph(style=f"Heading {min(block.level, 9)}"), block.text)
        case Kind.LIST_ITEM:
            kind = "Number" if block.ordered else "Bullet"
            level = min(block.depth + 1, _MAX_LIST_LEVEL)
            style = f"List {kind}" + (f" {level}" if level > 1 else "")
            _runs(out.add_paragraph(style=style), block.text)
        case Kind.TABLE:
            _table(out, block)
        case Kind.QUOTE:
            _runs(out.add_paragraph(style="Quote"), block.text)
        case Kind.CODE:
            _code(out.add_paragraph(), block.text)
        case Kind.PAGE_BREAK:
            out.add_page_break()
        case _:
            _runs(out.add_paragraph(), block.text)


def _code(par: Paragraph, text: str) -> None:
    run = par.add_run(text)
    run.font.name = _CODE_FONT


def _table(out: Any, block: Block) -> None:
    table = out.add_table(rows=len(block.rows), cols=len(block.rows[0]))
    table.style = "Light Grid Accent 1"
    for row, cells in zip(table.rows, block.rows, strict=True):
        for cell, text in zip(row.cells, cells, strict=True):
            _runs(cell.paragraphs[0], text)


def _runs(par: Paragraph, text: str) -> None:
    for run in parse_inline(text):
        if run.href:
            _hyperlink(par, run)
        else:
            _format(par.add_run(run.text), run)


def _format(target: DocxRun, run: Run) -> None:
    target.bold = run.bold or None
    target.italic = run.italic or None
    if run.code:
        target.font.name = _CODE_FONT


def _hyperlink(par: Paragraph, run: Run) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), par.part.relate_to(run.href, RT.HYPERLINK, is_external=True))
    element: Any = OxmlElement("w:r")
    link.append(element)
    par._p.append(link)  # pyright: ignore[reportPrivateUsage]
    styled = DocxRun(element, par)
    styled.text = run.text
    styled.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    styled.font.underline = True
    _format(styled, run)
